# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"C:\Users\26301\Documents\Codex\2026-08-23\c-users-26301-documents-codex-2026\outputs\改动结果汇总_2026-09-09\修改记录表_2026-09-09.xlsx"

headers = [
    "改动类型",
    "2026-09-09 上午\nRQ1 补充稳健性",
    "2026-09-09\nRQ2 clean3dim",
    "2026-09-09\nRQ3 clean3dim",
    "2026-09-09 下午\nWald / 图修正",
    "说明 / 关键数字",
]

rows = [
    ["RQ1 12控制主表", "√", "—", "—", "—",
     "表2第(3)列；func 0.0202 (p=0.014)、HCI 0.0130 (p=0.004)"],
    ["RQ1 平行趋势", "√ 事件研究", "—", "—", "√ 联合Wald",
     "原对角Wald改为聚类稳健协方差联合检验；func p=0.859、HCI p=0.955"],
    ["RQ1 安慰剂", "√ 500次城市置换", "—", "—", "—",
     "HCI p=0.038（双侧）、func p=0.074（双侧）"],
    ["RQ1 PSM-DID", "√ 1:1 / 1:4 / GDP四分位 / 有放回加权", "—", "—", "—",
     "方向与主结果一致，部分规格显著，定位为稳健性证据"],
    ["RQ1 图", "—", "—", "—", "√ 图例 + HCI 命名",
     "事件研究图与安慰剂图；BHCI 改为 HCI，红/灰线已加图例"],
    ["RQ2 三组异质性", "—", "√", "—", "—",
     "三维完整样本 4922/5299；High HCI 0.0220 (p=0.002, q=0.006)"],
    ["RQ2 组间对比与预趋势", "—", "√", "—", "√",
     "High-Low 0.0276 (p=0.0075)；预趋势 Low 0.455 / Mid 0.967 / High 0.436"],
    ["RQ3 维度×组", "—", "—", "√", "—",
     "High组物质/技能/动机完全脆弱均显著为正（HCI）"],
    ["RQ3 同分不同构成", "—", "—", "√", "—",
     "DVI3=0.6667；1|0.5|0.5 − 0.5|0.5|1 = 0.0519 (p=0.012, q=0.050)"],
    ["样本口径", "—", "√", "√", "—",
     "剔除任一维度缺失个案；0.25/0.75 伪档自动清除"],
    ["代码同步", "√ RQ1_supplement.R", "√ RQ2_three_group_clean3dim.R",
     "√ RQ3_clean3dim.R", "√ pre_trend_joint_wald.R / figures_RQ1.R",
     "所有代码均在“代码”文件夹"],
]

wb = Workbook()
ws = wb.active
ws.title = "修改记录"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True, size=10)
body_align = Alignment(vertical="center", wrap_text=True)
header_align = Alignment(vertical="center", horizontal="center", wrap_text=True)

for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = header_align
    c.border = border

for i, row in enumerate(rows, 2):
    for j, val in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=val)
        c.alignment = body_align
        c.border = border
        if j == 1:
            c.font = Font(bold=True, size=10)

widths = [26, 22, 16, 16, 20, 58]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 45
for i in range(2, len(rows) + 2):
    ws.row_dimensions[i].height = 46
ws.freeze_panes = "B2"

wb.save(OUT)
print("saved", OUT)
